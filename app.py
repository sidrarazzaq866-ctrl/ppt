import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import io
import re

st.set_page_config(page_title="Content to PowerPoint", page_icon="📊", layout="centered")

st.title("📊 Content to PowerPoint Converter")
st.write(
    "Paste your content below using a simple format — a title line starting with "
    "`#`, followed by bullet points starting with `-`. Each title starts a new slide. "
    "Click **Generate**, and you'll get a real, fully editable PowerPoint (.pptx) file — "
    "every text box can still be moved, resized, and edited normally in PowerPoint."
)

placeholder_text = """# Introduction
- Background of the topic
- Why it matters
- What this presentation covers

# Methodology
- Step one of the approach
- Step two of the approach
- Tools and materials used

# Results
- Key finding one
- Key finding two

# Conclusion
- Summary of findings
- Future work
"""

deck_title = st.text_input("Presentation title (shown on the first slide)", value="My Presentation")

content = st.text_area(
    "Slide content",
    height=320,
    placeholder=placeholder_text,
)

theme_color = st.color_picker("Accent color (slide titles)", value="#1F4E79")

filename = st.text_input("File name (without .pptx)", value="presentation")

generate_clicked = st.button("Generate PowerPoint", type="primary", use_container_width=True)


def parse_slides(text):
    """Split raw text into a list of (title, [bullets]) using '# ' and '- ' markers."""
    slides = []
    current_title = None
    current_bullets = []
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#"):
            if current_title is not None:
                slides.append((current_title, current_bullets))
            current_title = line.lstrip("#").strip()
            current_bullets = []
        elif line.startswith("-") or line.startswith("*"):
            current_bullets.append(line.lstrip("-*").strip())
        else:
            # A plain line under a title with no leading '-' is still treated as a bullet.
            current_bullets.append(line)
    if current_title is not None:
        slides.append((current_title, current_bullets))
    return slides


def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def build_pptx(deck_title, slides, color_hex):
    prs = Presentation()
    accent = hex_to_rgb(color_hex)

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = deck_title
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = accent
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = ""

    # Content slides
    bullet_layout = prs.slide_layouts[1]
    for title, bullets in slides:
        s = prs.slides.add_slide(bullet_layout)
        s.shapes.title.text = title
        s.shapes.title.text_frame.paragraphs[0].font.color.rgb = accent

        body = s.placeholders[1].text_frame
        body.clear()
        if not bullets:
            bullets = [""]
        for i, b in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = b
            p.level = 0
            p.font.size = Pt(20)

    return prs


if generate_clicked:
    if not content.strip():
        st.warning("Please paste some content first.")
    else:
        slides = parse_slides(content)
        if not slides:
            st.warning(
                "I couldn't find any slides in that text. Make sure each slide title "
                "starts with '#' on its own line."
            )
        else:
            prs = build_pptx(deck_title or "Presentation", slides, theme_color)
            buffer = io.BytesIO()
            prs.save(buffer)
            buffer.seek(0)

            st.success(f"Your presentation is ready — {len(slides) + 1} slides total.")
            st.download_button(
                "⬇️ Download PowerPoint (.pptx)",
                data=buffer,
                file_name=f"{filename or 'presentation'}.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True,
            )

            with st.expander("Preview slide structure"):
                st.write(f"**Slide 1 (Title):** {deck_title}")
                for i, (title, bullets) in enumerate(slides, start=2):
                    st.write(f"**Slide {i}:** {title}")
                    for b in bullets:
                        st.write(f"- {b}")

st.divider()
st.caption(
    "Tip: this creates a real editable PowerPoint, not an image. Once downloaded, "
    "you can change fonts, colors, layouts, or add images directly in PowerPoint."
)
